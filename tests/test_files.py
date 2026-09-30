"""create_document, the Markdown parser, and naming rules. No API key needed."""
import asyncio
import csv
import json
from pathlib import Path

import pytest

from bot import config
from bot.dispatcher import dispatch
from bot.tools import md_blocks, paths

SAMPLE_MD = """# Weekly Goals

Intro with **bold** and *italic* text.

## Tasks
- Ship the voice bot
- Write the report
1. First
2. Second

| Name | Capital |
|------|---------|
| India | New Delhi |

```
print("hi")
```
"""


def create(**args) -> dict:
    return json.loads(asyncio.run(dispatch("create_document", json.dumps(args))))


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "LOG_FILE", tmp_path / "test-bot.log")
    monkeypatch.setattr(config, "CONFIRM_DESTRUCTIVE", True)
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path / "output")


def assert_created(r: dict, ext: str) -> Path:
    assert r["ok"], r
    path = Path(r["path"])
    assert path.is_absolute() and path.suffix == f".{ext}"
    assert path.exists() and r["bytes"] > 0 and path.stat().st_size == r["bytes"]
    return path


# --- one test per format ---------------------------------------------------------

def test_md_prepends_title(tmp_path):
    path = assert_created(create(format="md", filename="notes", content="- a", title="My Notes", folder=str(tmp_path)), "md")
    assert path.read_text(encoding="utf-8").startswith("# My Notes\n")


def test_txt_strips_markdown(tmp_path):
    path = assert_created(create(format="txt", filename="plain", content="# Head\n**bold** line", folder=str(tmp_path)), "txt")
    assert path.read_text(encoding="utf-8") == "Head\nbold line\n"


def test_json_pretty_and_invalid(tmp_path):
    path = assert_created(create(format="json", filename="data", content='{"a": [1, 2]}', folder=str(tmp_path)), "json")
    assert json.loads(path.read_text(encoding="utf-8")) == {"a": [1, 2]}
    r = create(format="json", filename="bad", content="{nope", folder=str(tmp_path))
    assert r["ok"] is False and "JSON" in r["error"]


def test_html_from_markdown(tmp_path):
    path = assert_created(create(format="html", filename="page", content=SAMPLE_MD, folder=str(tmp_path)), "html")
    text = path.read_text(encoding="utf-8")
    assert "<h1>Weekly Goals</h1>" in text and "<li>Ship the voice bot</li>" in text
    assert "<strong>bold</strong>" in text and "<table>" in text


def test_csv_from_json_dicts_has_bom_and_header(tmp_path):
    rows = [{"country": "India", "capital": "New Delhi"}, {"country": "Côte d'Ivoire", "capital": "Yamoussoukro"}]
    path = assert_created(create(format="csv", filename="capitals", content=json.dumps(rows), folder=str(tmp_path)), "csv")
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    with path.open(encoding="utf-8-sig", newline="") as f:
        parsed = list(csv.reader(f))
    assert parsed == [["country", "capital"], ["India", "New Delhi"], ["Côte d'Ivoire", "Yamoussoukro"]]


def test_csv_from_text_and_markdown_table(tmp_path):
    path = assert_created(create(format="csv", filename="t1", content="a,b\n1,2", folder=str(tmp_path)), "csv")
    assert path.read_text(encoding="utf-8-sig").splitlines() == ["a,b", "1,2"]
    path = assert_created(create(format="csv", filename="t2", content="| a | b |\n|---|---|\n| 1 | 2 |", folder=str(tmp_path)), "csv")
    assert path.read_text(encoding="utf-8-sig").splitlines() == ["a,b", "1,2"]


def test_docx_has_heading_bullets_table(tmp_path):
    from docx import Document

    path = assert_created(create(format="docx", filename="goals", content=SAMPLE_MD, title="Weekly Goals", folder=str(tmp_path)), "docx")
    doc = Document(str(path))
    by_style = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Title", "Weekly Goals") in by_style
    assert ("Heading 2", "Tasks") in by_style
    assert ("List Bullet", "Ship the voice bot") in by_style
    assert ("List Number", "First") in by_style
    # The leading "# Weekly Goals" duplicates the title and is dropped.
    assert sum(1 for _, t in by_style if t == "Weekly Goals") == 1
    assert doc.tables[0].cell(1, 1).text == "New Delhi"
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "bold" in bold_runs


def test_pdf_valid_with_unicode(tmp_path):
    content = SAMPLE_MD + "\nUnicode: café, naïve, ₹ 500 — “quotes”.\n"
    path = assert_created(create(format="pdf", filename="report", content=content, title="Report", folder=str(tmp_path)), "pdf")
    assert path.read_bytes().startswith(b"%PDF")


def test_pdf_core_font_fallback(tmp_path, monkeypatch):
    from bot.tools import files

    monkeypatch.setattr(files, "_TTF_FAMILIES", [])
    monkeypatch.setattr(files, "_LINUX_MAC_FONTS", [])
    path = assert_created(create(format="pdf", filename="core", content="- café ₹ 5", folder=str(tmp_path)), "pdf")
    assert path.read_bytes().startswith(b"%PDF")


def test_xlsx_header_bold_and_numbers(tmp_path):
    from openpyxl import load_workbook

    content = json.dumps([["Country", "Population"], ["India", "1400000000"]])
    path = assert_created(create(format="xlsx", filename="pop", content=content, title="Data", folder=str(tmp_path)), "xlsx")
    ws = load_workbook(path).active
    assert ws.title == "Data" and ws["A1"].font.bold and ws["B2"].value == 1400000000
    assert ws.freeze_panes == "A2"


def test_format_aliases_and_unsupported(tmp_path):
    assert_created(create(format="word", filename="w", content="hi", folder=str(tmp_path)), "docx")
    r = create(format="pptx", filename="deck", content="x", folder=str(tmp_path))
    assert r["ok"] is False and "write_file" in r["error"]


# --- naming and overwrite rules --------------------------------------------------

def test_slug_and_auto_suffix(tmp_path):
    first = create(format="docx", filename="Weekly Goals!", content="- a", folder=str(tmp_path))
    second = create(format="docx", filename="Weekly Goals!", content="- b", folder=str(tmp_path))
    assert Path(first["path"]).name == "weekly-goals.docx"
    assert Path(second["path"]).name == "weekly-goals-1.docx"


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("Weekly Goals!", "weekly-goals"),
        ("report.doc", "report"),
        ('a\\b/c:d*e?f"g<h>i|j', "abcdefghij"),
        ("CON", "con-file"),
        ("  lots   of -- dashes  ", "lots-of-dashes"),
        ("x" * 200, "x" * 80),
    ],
)
def test_slugify(raw, expected):
    assert paths.slugify(raw) == expected


def test_slugify_empty_gets_dated_name():
    assert paths.slugify("!!!").startswith("document-")


def test_overwrite_needs_confirmation(tmp_path):
    create(format="md", filename="plan", content="old", folder=str(tmp_path))
    target = tmp_path / "plan.md"
    r = create(format="md", filename="plan", content="new", folder=str(tmp_path), overwrite=True)
    assert r["needs_confirmation"] is True
    assert target.read_text(encoding="utf-8") == "old\n"
    r = create(format="md", filename="plan", content="new", folder=str(tmp_path), overwrite=True, confirmed=True)
    assert r["ok"] and Path(r["path"]) == target.resolve()
    assert target.read_text(encoding="utf-8") == "new\n"


def test_file_open_elsewhere_is_friendly(tmp_path, monkeypatch):
    from bot.tools import files

    def locked(*_):
        raise PermissionError("locked")

    monkeypatch.setitem(files._WRITERS, "docx", locked)
    r = create(format="docx", filename="busy", content="x", folder=str(tmp_path))
    assert r["ok"] is False and "open in another program" in r["error"]


# --- folders ------------------------------------------------------------------------

def test_default_and_output_alias_use_output_dir(tmp_path):
    out = tmp_path / "output"
    assert Path(create(format="md", filename="a", content="x")["path"]).parent == out.resolve()
    assert Path(create(format="md", filename="b", content="x", folder="output")["path"]).parent == out.resolve()


def test_folder_aliases_map_to_home(monkeypatch, tmp_path):
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    (tmp_path / "Desktop").mkdir()
    assert paths.resolve_folder("desktop") == tmp_path / "Desktop"
    assert paths.resolve_folder("My Documents") == tmp_path / "Documents"
    assert paths.resolve_folder("downloads") == tmp_path / "Downloads"


def test_onedrive_desktop_fallback(monkeypatch, tmp_path):
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    (tmp_path / "OneDrive" / "Desktop").mkdir(parents=True)
    assert paths.resolve_folder("desktop") == tmp_path / "OneDrive" / "Desktop"


# --- parser -------------------------------------------------------------------------

def test_parse_blocks():
    kinds = [b["type"] for b in md_blocks.parse(SAMPLE_MD)]
    assert kinds == ["heading", "paragraph", "heading", "bullet", "bullet", "numbered", "numbered", "table", "code"]


def test_inline_runs():
    assert md_blocks.inline_runs("a **b** *c* `d`") == [
        ("a ", False, False), ("b", True, False), (" ", False, False),
        ("c", False, True), (" ", False, False), ("d", False, False),
    ]
